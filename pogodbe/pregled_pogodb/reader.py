"""Branje besedila iz datotek: PDF (besedilo ali sken), slike, Word."""
from __future__ import annotations

import hashlib
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Callable, Optional

from . import ocr

PDF_EXT = {".pdf"}
IMG_EXT = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp"}
DOC_EXT = {".docx"}
SUPPORTED = PDF_EXT | IMG_EXT | DOC_EXT


@dataclass
class ReadResult:
    text: str
    pages: int
    ocr_pages: int
    warning: str = ""


def file_hash(path: str) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def _clean(text: str) -> str:
    text = text.replace("\r", "")
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _read_pdf(path: str, progress: Optional[Callable[[str], None]]) -> ReadResult:
    import pymupdf as fitz

    pages, n_ocr, warn = [], 0, ""
    with fitz.open(path) as doc:
        total = len(doc)
        for i, page in enumerate(doc):
            txt = page.get_text("text") or ""
            if len(txt.strip()) < 40:  # sken brez besedilnega sloja
                if not ocr.available():
                    warn = "Sken brez OCR: Tesseract ni na voljo."
                else:
                    if progress:
                        progress(f"OCR stran {i + 1}/{total}")
                    pix = page.get_pixmap(matrix=fitz.Matrix(300 / 72, 300 / 72), colorspace=fitz.csGRAY)
                    txt = ocr.ocr_bytes(pix.tobytes("png"))
                    n_ocr += 1
            pages.append(txt)
    return ReadResult(_clean("\f".join(pages)), len(pages), n_ocr, warn)


def _read_image(path: str, progress: Optional[Callable[[str], None]]) -> ReadResult:
    if not ocr.available():
        return ReadResult("", 1, 0, "Slika brez OCR: Tesseract ni na voljo.")
    if progress:
        progress("OCR slike")
    return ReadResult(_clean(ocr.ocr_image(path)), 1, 1)


def _read_docx(path: str) -> ReadResult:
    from docx import Document

    d = Document(path)
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    return ReadResult(_clean("\n".join(parts)), 1, 0)


def read_file(path: str, progress: Optional[Callable[[str], None]] = None) -> ReadResult:
    ext = Path(path).suffix.lower()
    if ext in PDF_EXT:
        return _read_pdf(path, progress)
    if ext in IMG_EXT:
        return _read_image(path, progress)
    if ext in DOC_EXT:
        return _read_docx(path)
    raise ValueError(f"Nepodprta vrsta datoteke: {ext}")


def find_files(folder: str) -> list[str]:
    return sorted(str(p) for p in Path(folder).rglob("*")
                  if p.is_file() and p.suffix.lower() in SUPPORTED and not p.name.startswith("~$"))
