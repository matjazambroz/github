"""Izvoz seznama pogodb v Excel."""
from __future__ import annotations

from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .db import Store
from .extract import status_of

HEAD = ["Naslov", "Št. pogodbe", "Vrsta", "Stranka 1", "Stranka 2", "Sklenjena", "Velja od", "Velja do",
        "Status", "Samodejno podaljšanje", "Odpovedni rok", "Vrednost", "Valuta", "Obdobje", "DDV",
        "Plačilni rok (dni)", "Predmet", "Datoteka"]


def _d(s):
    return date.fromisoformat(s) if s else None


def export_xlsx(store: Store, path: str, ids: list[int] | None = None) -> int:
    wb = Workbook()
    ws = wb.active
    ws.title = "Pogodbe"
    ws.append(HEAD)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E79")
    n = 0
    for row in store.list():
        if ids is not None and row["id"] not in ids:
            continue
        r = store.get(row["id"])
        ws.append([r["title"], r["number"], r["doc_type"], r["party1"], r["party2"], _d(r["sign_date"]),
                   _d(r["start_date"]), _d(r["end_date"]),
                   status_of(_d(r["end_date"]), bool(r["indefinite"])), "da" if r["auto_renew"] else "ne",
                   r["notice_period"], r["value"], r["currency"], r["value_period"], r["vat_note"],
                   r["payment_term_days"], r["subject"], r["file_path"]])
        n += 1
    for col in range(1, len(HEAD) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 18
        for cell in ws[get_column_letter(col)][1:]:
            if isinstance(cell.value, date):
                cell.number_format = "DD.MM.YYYY"
            elif col == 12:
                cell.number_format = "#,##0.00"
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["Q"].width = 60
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(path)
    return n
